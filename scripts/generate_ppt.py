"""
scripts/generate_ppt.py

Generates a professional 5-slide PowerPoint presentation for the NDA Automated
Risk Assessment project using python-pptx.

Slides:
  1. Topic & Base Paper
  2. Objectives
  3. Model Trained, Why & Architecture
  4. Evaluation Parameters & Results
  5. Planned Improvements & Why

Usage:
    python scripts/generate_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─── Colour Palette ───────────────────────────────────────────────────────────
DARK_BG        = RGBColor(0x0D, 0x1B, 0x2A)
ACCENT_BLUE    = RGBColor(0x1A, 0x78, 0xC2)
ACCENT_CYAN    = RGBColor(0x00, 0xD4, 0xFF)
ACCENT_GREEN   = RGBColor(0x39, 0xD3, 0x53)
ACCENT_ORANGE  = RGBColor(0xFF, 0x6B, 0x35)
ACCENT_PURPLE  = RGBColor(0x9B, 0x59, 0xB6)
WHITE          = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY     = RGBColor(0xB0, 0xBE, 0xC5)
CARD_BG        = RGBColor(0x16, 0x2A, 0x3E)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)


def fill_shape(shape, rgb):
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def set_slide_bg(slide, rgb):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_text_box(slide, text, left, top, width, height,
                 font_size=18, bold=False, color=None,
                 align=PP_ALIGN.LEFT, italic=False, wrap=True):
    if color is None:
        color = WHITE
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox


def add_rect(slide, left, top, width, height, rgb):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    fill_shape(shape, rgb)
    shape.line.fill.background()
    return shape


def add_bullet_box(slide, items, left, top, width, height,
                   font_size=15, color=None):
    if color is None:
        color = WHITE
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = item
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
    return txBox


def add_accent_line(slide, left, top, width, rgb, thickness=4):
    line = slide.shapes.add_shape(1, left, top, width, Pt(thickness))
    fill_shape(line, rgb)
    line.line.fill.background()
    return line


# ═══════════════════════════════════════════════════════════════════════════════
#  SLIDE 1 – Topic & Base Paper
# ═══════════════════════════════════════════════════════════════════════════════
def slide1_topic(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), ACCENT_BLUE)
    add_rect(slide, Inches(0), Inches(0.12), Inches(13.33), Inches(0.04), ACCENT_CYAN)
    add_rect(slide, Inches(0), Inches(0), Inches(5.5), Inches(7.5), CARD_BG)
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), ACCENT_BLUE)
    add_rect(slide, Inches(0), Inches(0.12), Inches(13.33), Inches(0.04), ACCENT_CYAN)

    add_rect(slide, Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55), ACCENT_BLUE)
    add_text_box(slide, "01", Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55),
                 font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    add_text_box(slide, "NDA Automated Risk Assessment",
                 Inches(1.1), Inches(0.28), Inches(11.5), Inches(0.7),
                 font_size=34, bold=True, color=WHITE)
    add_accent_line(slide, Inches(1.1), Inches(1.0), Inches(8.5), ACCENT_CYAN, thickness=3)
    add_text_box(slide, "Topic & Base Paper",
                 Inches(1.1), Inches(1.08), Inches(6), Inches(0.4),
                 font_size=16, color=ACCENT_CYAN, bold=True)

    add_text_box(slide, "RESEARCH TOPIC",
                 Inches(0.3), Inches(1.65), Inches(4.8), Inches(0.42),
                 font_size=13, bold=True, color=ACCENT_CYAN)
    topic_text = (
        "Multi-Label Classification of NDA Clauses using Legal Language Models "
        "for Automated Document Risk Assessment & Signing Verification."
    )
    add_text_box(slide, topic_text, Inches(0.3), Inches(2.12), Inches(4.85), Inches(1.5),
                 font_size=13.5, color=WHITE, wrap=True)

    tags = [("Legal NLP", ACCENT_BLUE), ("Transformers", ACCENT_PURPLE), ("Contract AI", ACCENT_GREEN)]
    tx = Inches(0.3)
    for label, col in tags:
        add_rect(slide, tx, Inches(3.82), Inches(1.45), Inches(0.42), col)
        add_text_box(slide, label, tx, Inches(3.82), Inches(1.45), Inches(0.42),
                     font_size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        tx += Inches(1.6)

    add_text_box(slide, "DATASET",
                 Inches(0.3), Inches(4.45), Inches(4.8), Inches(0.38),
                 font_size=13, bold=True, color=ACCENT_CYAN)
    ds_bullets = [
        "  * Kleister-NDA Benchmark",
        "  * 20 NDA documents  |  717 clauses",
        "  * 14 legal taxonomy categories",
        "  * Document-disjoint split (0 overlap)",
    ]
    add_bullet_box(slide, ds_bullets, Inches(0.3), Inches(4.88), Inches(4.9), Inches(1.6),
                   font_size=13, color=LIGHT_GREY)

    add_text_box(slide, "BASE PAPER",
                 Inches(5.85), Inches(1.65), Inches(7.0), Inches(0.42),
                 font_size=13, bold=True, color=ACCENT_ORANGE)
    paper_title = (
        '"A Two-Stage Architecture for NDA Analysis:\n'
        ' LLM-based Segmentation and\n'
        ' Transformer-based Clause Classification"'
    )
    add_text_box(slide, paper_title, Inches(5.85), Inches(2.12), Inches(7.0), Inches(1.3),
                 font_size=14, italic=True, color=WHITE)

    details = [
        ("Dataset",       "Kleister-NDA  *  322 docs / 3,714 clauses"),
        ("Categories",    "14 approved legal taxonomy labels"),
        ("Architecture",  "Two-stage: LLM Segment -> Transformer Classify"),
        ("Gap Addressed", "Clause-level annotations not public;\n"
                          "Project builds pseudo-labels via Gemini 3-model ensemble"),
    ]
    ty = Inches(3.6)
    for label, val in details:
        add_rect(slide, Inches(5.85), ty, Inches(7.1), Inches(0.72), CARD_BG)
        add_text_box(slide, label.upper(),
                     Inches(6.0), ty + Pt(4), Inches(2.0), Inches(0.3),
                     font_size=10, bold=True, color=ACCENT_CYAN)
        add_text_box(slide, val, Inches(7.85), ty + Pt(4), Inches(4.8), Inches(0.62),
                     font_size=12, color=WHITE, wrap=True)
        ty += Inches(0.82)

    add_rect(slide, Inches(0), Inches(7.3), Inches(13.33), Inches(0.2), ACCENT_BLUE)
    add_text_box(slide, "NDA Risk Assessment  |  Research Presentation  |  2026",
                 Inches(0), Inches(7.28), Inches(13.33), Inches(0.22),
                 font_size=9, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
#  SLIDE 2 – Objectives
# ═══════════════════════════════════════════════════════════════════════════════
def slide2_objectives(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), ACCENT_CYAN)
    add_rect(slide, Inches(0), Inches(0.12), Inches(13.33), Inches(0.04), ACCENT_BLUE)

    add_rect(slide, Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55), ACCENT_CYAN)
    add_text_box(slide, "02", Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55),
                 font_size=16, bold=True, color=DARK_BG, align=PP_ALIGN.CENTER)
    add_text_box(slide, "Objectives",
                 Inches(1.1), Inches(0.28), Inches(10), Inches(0.7),
                 font_size=34, bold=True, color=WHITE)
    add_accent_line(slide, Inches(1.1), Inches(1.0), Inches(6), ACCENT_CYAN, thickness=3)
    add_text_box(slide, "What this project aims to achieve",
                 Inches(1.1), Inches(1.08), Inches(7), Inches(0.4),
                 font_size=15, color=LIGHT_GREY)

    objectives = [
        ("RQ1", "Architecture Benchmark",
         "Compare legal-domain (Legal-RoBERTa, Legal-BERT) vs. general-domain "
         "(DeBERTa-v3) transformers for NDA multi-label clause classification.",
         ACCENT_BLUE),
        ("RQ2", "Imbalance-Aware Loss Evaluation",
         "Quantify the effect of Class-Weighted BCE & Multi-Label Focal Loss on "
         "minority class recall and overall Macro F1 under extreme class imbalance.",
         ACCENT_CYAN),
        ("RQ3", "Pseudo-Label Confidence Trade-off",
         "Evaluate whether filtering Gemini LLM ensemble pseudo-labels by "
         "confidence level improves or hurts downstream classifier performance.",
         ACCENT_GREEN),
        ("RQ4", "Threshold Optimization & Risk Scoring",
         "Optimize decision boundary on validation data and deploy the best model "
         "as an end-to-end NDA PDF risk tool with signing recommendation.",
         ACCENT_ORANGE),
    ]
    positions = [
        (Inches(0.4),  Inches(1.65)),
        (Inches(6.95), Inches(1.65)),
        (Inches(0.4),  Inches(4.3)),
        (Inches(6.95), Inches(4.3)),
    ]
    col_w = Inches(5.9)
    for (left, top), (rq, title, desc, color) in zip(positions, objectives):
        add_rect(slide, left, top, col_w, Inches(2.4), CARD_BG)
        add_accent_line(slide, left, top, col_w, color, thickness=5)
        add_rect(slide, left + Inches(0.15), top + Inches(0.18), Inches(0.55), Inches(0.35), color)
        add_text_box(slide, rq,
                     left + Inches(0.15), top + Inches(0.18), Inches(0.55), Inches(0.35),
                     font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, title,
                     left + Inches(0.82), top + Inches(0.15), Inches(4.9), Inches(0.45),
                     font_size=14, bold=True, color=color)
        add_text_box(slide, desc,
                     left + Inches(0.15), top + Inches(0.72), Inches(5.6), Inches(1.55),
                     font_size=12.5, color=LIGHT_GREY, wrap=True)

    add_rect(slide, Inches(0), Inches(7.3), Inches(13.33), Inches(0.2), ACCENT_CYAN)
    add_text_box(slide, "NDA Risk Assessment  |  Research Presentation  |  2026",
                 Inches(0), Inches(7.28), Inches(13.33), Inches(0.22),
                 font_size=9, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
#  SLIDE 3 – Model, Why & Architecture
# ═══════════════════════════════════════════════════════════════════════════════
def slide3_model(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), ACCENT_GREEN)
    add_rect(slide, Inches(0), Inches(0.12), Inches(13.33), Inches(0.04), ACCENT_BLUE)

    add_rect(slide, Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55), ACCENT_GREEN)
    add_text_box(slide, "03", Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55),
                 font_size=16, bold=True, color=DARK_BG, align=PP_ALIGN.CENTER)
    add_text_box(slide, "Model Trained -- Why & Architecture",
                 Inches(1.1), Inches(0.28), Inches(11.5), Inches(0.7),
                 font_size=30, bold=True, color=WHITE)
    add_accent_line(slide, Inches(1.1), Inches(1.0), Inches(9), ACCENT_GREEN, thickness=3)

    add_rect(slide, Inches(0.4), Inches(1.12), Inches(12.5), Inches(0.6), ACCENT_BLUE)
    add_text_box(slide,
                 "Best Model: Legal-RoBERTa + Class-Weighted BCE (EXP-03 @ threshold 0.60)",
                 Inches(0.4), Inches(1.12), Inches(12.5), Inches(0.6),
                 font_size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    models = [
        ("Legal-RoBERTa", "saibo/legal-roberta-base",
         "Pre-trained on court cases, legislation & contracts.\nSuperior legal domain knowledge -- Best performer.",
         ACCENT_GREEN, "[WINNER]"),
        ("Legal-BERT", "nlpaueb/legal-bert-base-uncased",
         "Pre-trained on EU/US legal texts.\nReasonable legal context but lower than RoBERTa.",
         ACCENT_CYAN, "[2nd]"),
        ("DeBERTa-v3", "microsoft/deberta-v3-base",
         "Disentangled attention (general-domain).\nLacks legal pre-training -- lowest performance.",
         ACCENT_ORANGE, "[3rd]"),
    ]
    ty = Inches(1.92)
    for name, hf_id, desc, col, badge in models:
        add_rect(slide, Inches(0.4), ty, Inches(5.5), Inches(1.08), CARD_BG)
        add_accent_line(slide, Inches(0.4), ty, Inches(5.5), col, thickness=5)
        add_text_box(slide, f"{badge}  {name}",
                     Inches(0.65), ty + Pt(5), Inches(4.5), Inches(0.38),
                     font_size=14, bold=True, color=col)
        add_text_box(slide, hf_id,
                     Inches(0.65), ty + Inches(0.42), Inches(4.8), Inches(0.28),
                     font_size=10, italic=True, color=LIGHT_GREY)
        add_text_box(slide, desc,
                     Inches(0.65), ty + Inches(0.65), Inches(5.0), Inches(0.52),
                     font_size=11.5, color=WHITE, wrap=True)
        ty += Inches(1.18)

    add_text_box(slide, "Pipeline Architecture",
                 Inches(6.3), Inches(1.92), Inches(6.5), Inches(0.42),
                 font_size=14, bold=True, color=ACCENT_CYAN)

    steps = [
        ("PDF Input",           "PyMuPDF text extraction",              ACCENT_BLUE),
        ("Clause Segmentation", "Rule-based + LLM-assisted splitter",   ACCENT_CYAN),
        ("Tokenization",        "AutoTokenizer (max 512 tokens)",        ACCENT_GREEN),
        ("Legal-RoBERTa Enc.",  "768-dim contextual embeddings",         ACCENT_ORANGE),
        ("Multi-Label Head",    "14-unit sigmoid + Class-Wt. BCE",       ACCENT_PURPLE),
        ("Risk Scoring",        "Threshold 0.60 -> Risk Index (0-100)",  ACCENT_CYAN),
    ]
    ty = Inches(2.42)
    for label, sub, col in steps:
        add_rect(slide, Inches(6.3), ty, Inches(6.6), Inches(0.65), CARD_BG)
        add_rect(slide, Inches(6.3), ty, Inches(0.55), Inches(0.65), col)
        add_text_box(slide, label, Inches(6.95), ty + Pt(3), Inches(3.2), Inches(0.35),
                     font_size=13, bold=True, color=WHITE)
        add_text_box(slide, sub, Inches(6.95), ty + Inches(0.35), Inches(5.8), Inches(0.28),
                     font_size=11, color=LIGHT_GREY)
        ty += Inches(0.78)

    add_rect(slide, Inches(0), Inches(7.3), Inches(13.33), Inches(0.2), ACCENT_GREEN)
    add_text_box(slide, "NDA Risk Assessment  |  Research Presentation  |  2026",
                 Inches(0), Inches(7.28), Inches(13.33), Inches(0.22),
                 font_size=9, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
#  SLIDE 4 – Evaluation Parameters & Results
# ═══════════════════════════════════════════════════════════════════════════════
def slide4_results(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), ACCENT_ORANGE)
    add_rect(slide, Inches(0), Inches(0.12), Inches(13.33), Inches(0.04), ACCENT_BLUE)

    add_rect(slide, Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55), ACCENT_ORANGE)
    add_text_box(slide, "04", Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55),
                 font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, "Evaluation Parameters & Results",
                 Inches(1.1), Inches(0.28), Inches(11.5), Inches(0.7),
                 font_size=30, bold=True, color=WHITE)
    add_accent_line(slide, Inches(1.1), Inches(1.0), Inches(8.5), ACCENT_ORANGE, thickness=3)

    kpis = [
        ("Macro F1 @ 0.60",     "0.4295", ACCENT_GREEN,  "+0.0620"),
        ("Minority F1 @ 0.60",  "0.3853", ACCENT_CYAN,   "+0.1047"),
        ("Hamming Loss @ 0.60", "0.1283", ACCENT_ORANGE, "-46.9% FP"),
        ("MCC @ 0.60",          "0.3885", ACCENT_PURPLE, "+0.0430"),
    ]
    kx = Inches(0.4)
    for label, val, col, delta in kpis:
        add_rect(slide, kx, Inches(1.15), Inches(2.9), Inches(1.3), CARD_BG)
        add_accent_line(slide, kx, Inches(1.15), Inches(2.9), col, thickness=4)
        add_text_box(slide, val, kx, Inches(1.32), Inches(2.9), Inches(0.65),
                     font_size=26, bold=True, color=col, align=PP_ALIGN.CENTER)
        add_text_box(slide, delta,
                     kx + Inches(1.8), Inches(1.32), Inches(0.95), Inches(0.35),
                     font_size=11, bold=True, color=ACCENT_GREEN)
        add_text_box(slide, label, kx, Inches(1.95), Inches(2.9), Inches(0.5),
                     font_size=11, color=LIGHT_GREY, align=PP_ALIGN.CENTER)
        kx += Inches(3.18)

    add_text_box(slide, "9-Experiment Matrix (default threshold 0.50)",
                 Inches(0.4), Inches(2.65), Inches(7.0), Inches(0.38),
                 font_size=13, bold=True, color=ACCENT_ORANGE)

    headers = ["Rank", "Exp ID", "Architecture", "Loss Function", "Test Macro F1"]
    widths  = [Inches(0.55), Inches(0.7), Inches(2.2), Inches(2.55), Inches(1.55)]
    hx = Inches(0.4)
    for h, w in zip(headers, widths):
        add_rect(slide, hx, Inches(3.08), w, Inches(0.38), ACCENT_BLUE)
        add_text_box(slide, h, hx, Inches(3.08), w, Inches(0.38),
                     font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        hx += w

    rows = [
        ("1", "EXP-03", "Legal-RoBERTa", "Class-Weighted BCE", "0.3675  [BEST]"),
        ("2", "EXP-06", "Legal-BERT",    "Class-Weighted BCE", "0.3047"),
        ("3", "EXP-09", "DeBERTa-v3",    "Class-Weighted BCE", "0.2739"),
        ("4", "EXP-01", "Legal-RoBERTa", "BCE (Standard)",     "0.0492"),
        ("—", "EXP-05", "Legal-BERT",    "Focal Loss",         "0.0000  [FAIL]"),
    ]
    ty = Inches(3.46)
    for i, row in enumerate(rows):
        bg = CARD_BG if i % 2 == 0 else DARK_BG
        rx = Inches(0.4)
        for cell, w in zip(row, widths):
            add_rect(slide, rx, ty, w, Inches(0.38), bg)
            col = ACCENT_GREEN if "[BEST]" in cell else (RGBColor(0xFF, 0x44, 0x44) if "[FAIL]" in cell else WHITE)
            add_text_box(slide, cell, rx, ty, w, Inches(0.38),
                         font_size=11, color=col, align=PP_ALIGN.CENTER)
            rx += w
        ty += Inches(0.38)

    add_text_box(slide, "Threshold Sweep (Validation Macro F1)",
                 Inches(7.8), Inches(2.65), Inches(5.2), Inches(0.38),
                 font_size=13, bold=True, color=ACCENT_CYAN)

    thresholds = [
        ("0.30", 0.2366, LIGHT_GREY),
        ("0.40", 0.3086, LIGHT_GREY),
        ("0.50", 0.3875, LIGHT_GREY),
        ("0.60 [OPT]", 0.4306, ACCENT_GREEN),
        ("0.70", 0.3917, LIGHT_GREY),
    ]
    ty = Inches(3.08)
    bar_max_w = Inches(3.5)
    for thr, val, col in thresholds:
        bar_w = bar_max_w * val
        add_rect(slide, Inches(7.8), ty, bar_w, Inches(0.35), col)
        add_text_box(slide, thr, Inches(7.8), ty, Inches(0.9), Inches(0.35),
                     font_size=11, bold=True,
                     color=DARK_BG if col == ACCENT_GREEN else LIGHT_GREY)
        add_text_box(slide, f"{val:.4f}",
                     Inches(7.8) + bar_w + Inches(0.08), ty, Inches(0.8), Inches(0.35),
                     font_size=11, color=col)
        ty += Inches(0.43)

    add_rect(slide, Inches(0.4), Inches(5.55), Inches(12.5), Inches(0.7), CARD_BG)
    add_text_box(slide,
                 "Pseudo-Label Ablation:  ALL-VALID (344 clauses) F1=0.3477  vs  "
                 "HIGH+MEDIUM (201) F1=0.2748  vs  HIGH-ONLY (161) F1=0.2485\n"
                 "=> More training volume beats confidence filtering under extreme class imbalance.",
                 Inches(0.55), Inches(5.58), Inches(12.2), Inches(0.65),
                 font_size=11.5, color=WHITE, wrap=True)

    cats = [
        ("Additional Info   F1=0.6494", ACCENT_GREEN),
        ("Liability/Damages F1=0.5957", ACCENT_CYAN),
        ("IP Rights         F1=0.5714", ACCENT_BLUE),
    ]
    cx = Inches(0.4)
    for cat, col in cats:
        add_rect(slide, cx, Inches(6.35), Inches(3.85), Inches(0.6), col)
        add_text_box(slide, cat, cx + Inches(0.1), Inches(6.35), Inches(3.5), Inches(0.6),
                     font_size=13, bold=True, color=WHITE)
        cx += Inches(4.15)

    add_rect(slide, Inches(0), Inches(7.3), Inches(13.33), Inches(0.2), ACCENT_ORANGE)
    add_text_box(slide, "NDA Risk Assessment  |  Research Presentation  |  2026",
                 Inches(0), Inches(7.28), Inches(13.33), Inches(0.22),
                 font_size=9, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
#  SLIDE 5 – Planned Improvements
# ═══════════════════════════════════════════════════════════════════════════════
def slide5_improvements(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), ACCENT_PURPLE)
    add_rect(slide, Inches(0), Inches(0.12), Inches(13.33), Inches(0.04), ACCENT_CYAN)

    add_rect(slide, Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55), ACCENT_PURPLE)
    add_text_box(slide, "05", Inches(0.3), Inches(0.35), Inches(0.55), Inches(0.55),
                 font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, "Planned Improvements & Why",
                 Inches(1.1), Inches(0.28), Inches(11.5), Inches(0.7),
                 font_size=32, bold=True, color=WHITE)
    add_accent_line(slide, Inches(1.1), Inches(1.0), Inches(8.5), ACCENT_PURPLE, thickness=3)
    add_text_box(slide, "Future roadmap to push Macro F1 beyond 0.60+",
                 Inches(1.1), Inches(1.08), Inches(8.5), Inches(0.38),
                 font_size=14, color=LIGHT_GREY)

    improvements = [
        ("1", "Human-Verified Annotations", ACCENT_BLUE,
         "Replace Gemini pseudo-labels with expert-annotated ground truth to eliminate "
         "label noise (current 32% disagreement rate) and get reliable F1 estimates."),
        ("2", "Larger & Diverse NDA Corpus", ACCENT_CYAN,
         "Scale from 20 -> 100+ NDA docs across industries. Minority classes like "
         "'NDA Type' & 'Purpose' have <8 clauses; more data will balance distribution."),
        ("3", "Contrastive Pre-training", ACCENT_GREEN,
         "Fine-tune with contrastive loss on legal clause pairs to reduce multi-label "
         "confusion between semantically similar categories (e.g. Confidentiality vs. Disclosure)."),
        ("4", "LLM-Based Smart Segmentation", ACCENT_ORANGE,
         "Replace heuristic segmenter with Gemma/Mistral LLM segmenter to produce "
         "cleaner clause boundaries, improving tokenizer coverage & classification quality."),
        ("5", "RAG-Powered Clause Explanation", ACCENT_PURPLE,
         "Add retrieval-augmented generation to explain each risk clause with legal "
         "precedents. Lawyers need interpretable summaries, not just category labels."),
        ("6", "GUI / API Deployment", RGBColor(0xFF, 0x6B, 0x35),
         "Build FastAPI backend + Streamlit/React frontend to make NDA risk assessment "
         "accessible to legal professionals without CLI / technical knowledge."),
    ]
    positions = [
        (Inches(0.4),  Inches(1.65)),
        (Inches(4.7),  Inches(1.65)),
        (Inches(9.0),  Inches(1.65)),
        (Inches(0.4),  Inches(4.18)),
        (Inches(4.7),  Inches(4.18)),
        (Inches(9.0),  Inches(4.18)),
    ]
    card_w = Inches(4.1)
    card_h = Inches(2.3)
    for (left, top), (num, title, col, desc) in zip(positions, improvements):
        add_rect(slide, left, top, card_w, card_h, CARD_BG)
        add_accent_line(slide, left, top, card_w, col, thickness=5)
        add_rect(slide, left + Inches(0.12), top + Inches(0.12),
                 Inches(0.42), Inches(0.42), col)
        add_text_box(slide, num,
                     left + Inches(0.12), top + Inches(0.12),
                     Inches(0.42), Inches(0.42),
                     font_size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, title,
                     left + Inches(0.65), top + Inches(0.12),
                     Inches(3.3), Inches(0.42),
                     font_size=12.5, bold=True, color=col)
        add_text_box(slide, desc,
                     left + Inches(0.12), top + Inches(0.65),
                     Inches(3.85), Inches(1.55),
                     font_size=11, color=LIGHT_GREY, wrap=True)

    add_rect(slide, Inches(0), Inches(7.3), Inches(13.33), Inches(0.2), ACCENT_PURPLE)
    add_text_box(slide, "NDA Risk Assessment  |  Research Presentation  |  2026",
                 Inches(0), Inches(7.28), Inches(13.33), Inches(0.22),
                 font_size=9, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════════════════════════
def main():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    print("Building Slide 1 - Topic & Base Paper...")
    slide1_topic(prs)

    print("Building Slide 2 - Objectives...")
    slide2_objectives(prs)

    print("Building Slide 3 - Model, Why & Architecture...")
    slide3_model(prs)

    print("Building Slide 4 - Evaluation Parameters & Results...")
    slide4_results(prs)

    print("Building Slide 5 - Planned Improvements & Why...")
    slide5_improvements(prs)

    out_path = "NDA_Risk_Assessment_Presentation.pptx"
    prs.save(out_path)
    print(f"\nSaved: {out_path}")


if __name__ == "__main__":
    main()
